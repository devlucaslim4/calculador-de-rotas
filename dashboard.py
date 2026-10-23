"""Componentes Streamlit do dashboard de análise."""

from __future__ import annotations

import logging
import math
from io import BytesIO

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import pydeck as pdk
import streamlit as st
from openpyxl import load_workbook

from analysis_engine import (
    AnalysisData,
    analysis_workbook,
    apply_filters,
    build_audit,
    calculate_metrics,
    normalize_header,
    prepare_analysis,
)
from report_export import analysis_pdf
from route_processor import RouteGeometry, RouteGeometryError, fetch_route_geometry, parse_coordinate

COLOR = "#3B82F6"
WEEKDAYS = {0: "Segunda", 1: "Terça", 2: "Quarta", 3: "Quinta", 4: "Sexta", 5: "Sábado", 6: "Domingo"}
DASHBOARD_FILTER_FIELDS = ("usuario", "unidade", "regiao", "motivo", "status")


@st.cache_data(ttl=86_400, show_spinner=False)
def _cached_route_geometry(origin: str, destination: str):
    """Mantém o traçado escolhido em cache sem atrasar a importação em lote."""

    return fetch_route_geometry(origin, destination)


def dataframe_from_excel(file_bytes: bytes) -> pd.DataFrame:
    frame = pd.read_excel(BytesIO(file_bytes), engine="openpyxl")
    workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    worksheet = workbook.active
    headers = {cell.column: cell.value for cell in worksheet[1]}
    for column_index, header in headers.items():
        if header not in frame.columns:
            continue
        dataframe_index = frame.columns.get_loc(header)
        route_cells = next(
            worksheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
                max_row=worksheet.max_row,
            )
        )
        for row_index, cell in enumerate(route_cells):
            if cell.hyperlink and cell.hyperlink.target:
                frame.iat[row_index, dataframe_index] = cell.hyperlink.target
    return frame


def render_dashboard(dataframe: pd.DataFrame, original_name: str, light_mode: bool = False) -> None:
    analysis = prepare_analysis(dataframe)
    for warning in analysis.warnings:
        st.warning(warning)

    st.subheader("Filtros")
    selections = _render_filters(analysis)

    filtered = apply_filters(analysis, selections)
    audit_all = build_audit(analysis, divergence_limit=20, duration_limit=8)
    audit = audit_all[audit_all["Índice da linha"].isin(filtered.index)].copy()
    metrics = calculate_metrics(analysis, filtered, set(audit_all["Índice da linha"].tolist()))

    if filtered.empty:
        st.info("Nenhum registro corresponde aos filtros selecionados.")
        return

    st.caption(f"Exibindo {len(filtered):,.0f} de {len(analysis.data):,.0f} registros".replace(",", "."))
    _render_metrics(metrics)
    routes = analysis.unique_routes(filtered)
    _render_downloads(analysis, filtered, audit, metrics, original_name)
    overview_tab, data_tab, audit_tab = st.tabs(["Visão geral", "Dados detalhados", "Auditoria"])
    with overview_tab:
        _render_charts(analysis, routes, 10, light_mode)
    with data_tab:
        _render_data_table(analysis, filtered)
    with audit_tab:
        _render_audit(audit)


def render_route_map(dataframe: pd.DataFrame, light_mode: bool = False) -> None:
    """Mostra uma rota rodoviária por vez em perspectiva 3D."""

    st.header("Mapa 3D das rotas")
    st.caption(
        "Filtre por usuário, escolha uma rota e carregue o trajeto. "
        "Arraste para mover, use a roda para aplicar zoom e Ctrl + arraste para inclinar ou girar."
    )
    analysis = prepare_analysis(dataframe)
    initial_column = analysis.columns.get("coord_inicial")
    final_column = analysis.columns.get("coord_final")
    if not initial_column or not final_column:
        st.warning("As colunas de coordenada inicial e final não foram encontradas.")
        return

    routes = analysis.unique_routes().copy()
    routes["__map_origin"] = routes[initial_column].map(parse_coordinate)
    routes["__map_destination"] = routes[final_column].map(parse_coordinate)
    routes = routes.loc[
        routes["__map_origin"].notna() & routes["__map_destination"].notna()
    ].copy()
    if routes.empty:
        st.info("Nenhuma rota com as duas coordenadas válidas está disponível para o mapa.")
        return

    user_column = analysis.columns.get("usuario")
    selector_left, selector_right = st.columns([1, 2])
    if user_column:
        users = sorted(
            routes[user_column]
            .dropna()
            .astype(str)
            .str.strip()
            .loc[lambda values: values.ne("")]
            .unique()
            .tolist()
        )
        selected_users = selector_left.multiselect(
            "Filtrar por usuário",
            users,
            key="map_users",
            placeholder="Todos os usuários",
        )
        if selected_users:
            routes = routes.loc[routes[user_column].astype(str).str.strip().isin(selected_users)]
    else:
        selector_left.info("Coluna de usuário não identificada.")

    if routes.empty:
        st.info("Nenhuma rota corresponde aos usuários selecionados.")
        return
    routes = routes.reset_index(drop=False).rename(columns={"index": "__source_index"})
    route_positions = list(range(len(routes)))
    selected_position = selector_right.selectbox(
        "Escolha uma rota",
        route_positions,
        format_func=lambda position: _map_route_label(routes.iloc[position], analysis),
        key="map_route_position",
    )
    row = routes.iloc[int(selected_position)]
    origin = row["__map_origin"]
    destination = row["__map_destination"]

    distance = row.get("__distancia_gps_valida")
    odometer = row.get("__distancia_hodometro")
    difference = row.get("__diferenca_percentual")
    metric_columns = st.columns(4)
    metric_columns[0].metric("Usuário", _map_cell(row, user_column, "Não informado"))
    metric_columns[1].metric("Distância GPS", _map_km(distance))
    metric_columns[2].metric("Distância hodômetro", _map_km(odometer))
    metric_columns[3].metric(
        "Diferença",
        "Não disponível" if pd.isna(difference) else f"{float(difference):.2f}%".replace(".", ","),
    )

    signature = (
        f"{origin[0]:.7f},{origin[1]:.7f}",
        f"{destination[0]:.7f},{destination[1]:.7f}",
    )
    if st.button(
        "Carregar experiência da rota",
        type="primary",
        width="stretch",
        key="load_3d_route",
    ):
        st.session_state["loaded_3d_route"] = signature

    if st.session_state.get("loaded_3d_route") != signature:
        st.info("Selecione a rota desejada e clique no botão para carregar a viagem interativa.")
        return

    try:
        with st.spinner("Carregando o traçado rodoviário detalhado…"):
            geometry = _cached_route_geometry(*signature)
    except RouteGeometryError as exc:
        st.error(str(exc))
        return
    except Exception:
        logging.exception("Falha inesperada ao montar o mapa 3D")
        st.error("Não foi possível montar o mapa desta rota agora.")
        return

    path = [list(point) for point in geometry.path]
    label = _map_route_label(row, analysis)
    view_mode = st.segmented_control(
        "Modo de visualização",
        ("Viagem animada", "Explorar em 3D"),
        default="Viagem animada",
        key="route_map_view_mode",
        width="stretch",
    )
    if view_mode == "Viagem animada":
        st.markdown(
            '<div class="result-callout"><strong>Viagem pronta.</strong> '
            'Use os controles no mapa para iniciar, pausar ou reiniciar o percurso.</div>',
            unsafe_allow_html=True,
        )
        figure = _animated_route_figure(path, label, geometry, light_mode)
        st.plotly_chart(
            figure,
            width="stretch",
            config={"displaylogo": False, "displayModeBar": False, "scrollZoom": True},
            key=f"animated_route_{signature[0]}_{signature[1]}",
        )
    else:
        _render_explorable_3d_map(path, label, geometry, origin, destination, light_mode, signature)
    st.caption(
        f"Traçado detalhado: {geometry.distance_km:.2f} km • tempo estimado: "
        f"{geometry.duration_minutes:.0f} min • verde = início • vermelho = destino"
    )


def _render_explorable_3d_map(
    path: list[list[float]],
    label: str,
    geometry: RouteGeometry,
    origin: tuple[float, float],
    destination: tuple[float, float],
    light_mode: bool,
    signature: tuple[str, str],
) -> None:
    marker_radius = max(35.0, min(650.0, geometry.distance_km * 12))
    marker_height = max(180.0, min(4_500.0, geometry.distance_km * 70))
    route_data = [{
        "path": path,
        "color": [59, 130, 246, 230],
        "label": label,
        "details": f"{geometry.distance_km:.2f} km • {geometry.duration_minutes:.0f} min",
    }]
    marker_data = [
        {
            "position": path[0],
            "color": [16, 185, 129, 225],
            "elevation": marker_height,
            "label": "Início",
            "details": f"{origin[0]:.6f}, {origin[1]:.6f}",
        },
        {
            "position": path[-1],
            "color": [244, 63, 94, 225],
            "elevation": marker_height,
            "label": "Destino",
            "details": f"{destination[0]:.6f}, {destination[1]:.6f}",
        },
    ]
    layers = [
        pdk.Layer(
            "PathLayer",
            route_data,
            get_path="path",
            get_color="color",
            get_width=6,
            width_units="pixels",
            width_min_pixels=4,
            pickable=True,
        ),
        pdk.Layer(
            "ColumnLayer",
            marker_data,
            get_position="position",
            get_fill_color="color",
            get_elevation="elevation",
            radius=marker_radius,
            disk_resolution=32,
            extruded=True,
            pickable=True,
        ),
    ]
    view_state = _map_view_state(path)
    deck = pdk.Deck(
        layers=layers,
        initial_view_state=view_state,
        map_style=pdk.map_styles.LIGHT if light_mode else pdk.map_styles.DARK,
        map_provider="carto",
        tooltip={
            "html": "<b>{label}</b><br>{details}",
            "style": {
                "backgroundColor": "#0f172a",
                "color": "#f8fafc",
                "fontSize": "13px",
            },
        },
    )
    st.pydeck_chart(deck, width="stretch", height=620, key=f"route_map_{signature[0]}_{signature[1]}")


def _animated_route_figure(
    path: list[list[float]],
    label: str,
    geometry: RouteGeometry,
    light_mode: bool,
) -> go.Figure:
    # Menos quadros mantêm o movimento suave e reduzem bastante o JSON enviado
    # ao navegador em cada nova seleção de rota.
    sampled = _sample_map_path(path, max_points=60)
    longitudes = [point[0] for point in sampled]
    latitudes = [point[1] for point in sampled]
    view_state = _map_view_state(path)
    muted_route = "rgba(96, 165, 250, 0.32)" if not light_mode else "rgba(37, 99, 235, 0.28)"
    panel_color = "#0f172a" if not light_mode else "#ffffff"
    text_color = "#f8fafc" if not light_mode else "#172033"

    figure = go.Figure()
    figure.add_trace(
        go.Scattermap(
            lon=longitudes,
            lat=latitudes,
            mode="lines",
            line={"width": 11, "color": muted_route},
            hoverinfo="skip",
            name="Trajeto completo",
        )
    )
    figure.add_trace(
        go.Scattermap(
            lon=[longitudes[0]],
            lat=[latitudes[0]],
            mode="lines",
            line={"width": 7, "color": "#3b82f6"},
            hoverinfo="skip",
            name="Trecho percorrido",
        )
    )
    figure.add_trace(
        go.Scattermap(
            lon=[longitudes[0], longitudes[-1]],
            lat=[latitudes[0], latitudes[-1]],
            mode="markers",
            text=["INÍCIO", "DESTINO"],
            marker={"size": [17, 19], "color": ["#10b981", "#f43f5e"]},
            hovertemplate="<b>%{text}</b><br>%{lat:.6f}, %{lon:.6f}<extra></extra>",
            name="Pontos da rota",
        )
    )
    figure.add_trace(
        go.Scattermap(
            lon=[longitudes[0]],
            lat=[latitudes[0]],
            mode="markers+text",
            text=["CARRO"],
            textposition="top center",
            textfont={"size": 12, "color": "#fbbf24"},
            marker={
                "size": 20,
                "color": "#f59e0b",
                "opacity": 1,
                "allowoverlap": True,
            },
            hovertemplate=f"<b>{label}</b><br>Veículo em movimento<extra></extra>",
            name="Veículo",
        )
    )

    frames: list[go.Frame] = []
    for index in range(len(sampled)):
        progress = int(round(index / max(len(sampled) - 1, 1) * 100))
        frames.append(
            go.Frame(
                name=str(index),
                data=[
                    go.Scattermap(
                        lon=longitudes[: index + 1],
                        lat=latitudes[: index + 1],
                        mode="lines",
                        line={"width": 7, "color": "#3b82f6"},
                        hoverinfo="skip",
                    ),
                    go.Scattermap(
                        lon=[longitudes[index]],
                        lat=[latitudes[index]],
                        mode="markers+text",
                        text=["CARRO"],
                        textposition="top center",
                        textfont={"size": 12, "color": "#fbbf24"},
                        marker={
                            "size": 20,
                            "color": "#f59e0b",
                            "opacity": 1,
                            "allowoverlap": True,
                        },
                        customdata=[progress],
                        hovertemplate="<b>Progresso: %{customdata}%</b><extra></extra>",
                    ),
                ],
                traces=[1, 3],
            )
        )
    figure.frames = frames
    figure.update_layout(
        height=650,
        margin={"l": 0, "r": 0, "t": 58, "b": 0},
        paper_bgcolor="rgba(0,0,0,0)",
        font={"color": text_color},
        showlegend=False,
        title={
            "text": f"<b>{label}</b><br><sup>{geometry.distance_km:.2f} km • {geometry.duration_minutes:.0f} min</sup>",
            "x": 0.02,
            "xanchor": "left",
        },
        map={
            "style": "carto-positron" if light_mode else "carto-darkmatter",
            "center": {"lon": view_state.longitude, "lat": view_state.latitude},
            "zoom": view_state.zoom,
            "pitch": 55,
            "bearing": view_state.bearing,
        },
        updatemenus=[{
            "type": "buttons",
            "direction": "left",
            "showactive": False,
            "x": 0.02,
            "y": 0.03,
            "xanchor": "left",
            "yanchor": "bottom",
            "bgcolor": panel_color,
            "bordercolor": "#3b82f6",
            "font": {"color": text_color},
            "buttons": [
                {
                    "label": "Iniciar viagem",
                    "method": "animate",
                    "args": [None, {
                        "frame": {"duration": 55, "redraw": True},
                        "transition": {"duration": 15},
                        "fromcurrent": True,
                        "mode": "immediate",
                    }],
                },
                {
                    "label": "Pausar",
                    "method": "animate",
                    "args": [[None], {
                        "frame": {"duration": 0, "redraw": False},
                        "transition": {"duration": 0},
                        "mode": "immediate",
                    }],
                },
                {
                    "label": "Reiniciar",
                    "method": "animate",
                    "args": [["0"], {
                        "frame": {"duration": 0, "redraw": True},
                        "transition": {"duration": 0},
                        "mode": "immediate",
                    }],
                },
            ],
        }],
    )
    return figure


def _map_route_label(row: pd.Series, analysis: AnalysisData) -> str:
    id_value = _map_cell(row, analysis.columns.get("id"), f"Linha {int(row.name) + 1}")
    user = _map_cell(row, analysis.columns.get("usuario"), "Usuário não informado")
    unit = _map_cell(row, analysis.columns.get("unidade"), "Unidade não informada")
    return f"{id_value} • {user} • {unit}"


def _map_cell(row: pd.Series, column: str | None, fallback: str) -> str:
    if not column:
        return fallback
    value = row.get(column)
    if value is None or pd.isna(value) or not str(value).strip():
        return fallback
    return str(value).strip()


def _map_km(value: object) -> str:
    if value is None or pd.isna(value):
        return "Não disponível"
    return f"{float(value):,.2f} km".replace(",", "X").replace(".", ",").replace("X", ".")


def _map_view_state(path: list[list[float]]) -> pdk.ViewState:
    longitudes = [point[0] for point in path]
    latitudes = [point[1] for point in path]
    longitude_span = max(longitudes) - min(longitudes)
    latitude_span = max(latitudes) - min(latitudes)
    span = max(longitude_span, latitude_span, 0.0005)
    zoom = max(3.0, min(16.0, math.log2(360 / span) - 1.65))
    delta_longitude = path[-1][0] - path[0][0]
    delta_latitude = path[-1][1] - path[0][1]
    bearing = math.degrees(math.atan2(delta_longitude, delta_latitude))
    return pdk.ViewState(
        longitude=(min(longitudes) + max(longitudes)) / 2,
        latitude=(min(latitudes) + max(latitudes)) / 2,
        zoom=zoom,
        pitch=55,
        bearing=bearing,
        min_zoom=2,
        max_zoom=19,
    )


def _sample_map_path(path: list[list[float]], max_points: int = 90) -> list[list[float]]:
    """Reamostra o caminho para o veículo avançar com velocidade visual uniforme."""

    if len(path) <= 2 or len(path) <= max_points:
        return [list(point) for point in path]
    reference_latitude = math.radians(sum(point[1] for point in path) / len(path))
    longitude_scale = max(math.cos(reference_latitude), 0.01)
    cumulative = [0.0]
    for previous, current in zip(path, path[1:]):
        delta_longitude = (current[0] - previous[0]) * longitude_scale
        delta_latitude = current[1] - previous[1]
        cumulative.append(cumulative[-1] + math.hypot(delta_longitude, delta_latitude))
    total = cumulative[-1]
    if total <= 0:
        return [list(path[0]), list(path[-1])]

    sampled: list[list[float]] = []
    segment = 1
    for position in range(max_points):
        target = total * position / (max_points - 1)
        while segment < len(cumulative) - 1 and cumulative[segment] < target:
            segment += 1
        start_distance = cumulative[segment - 1]
        end_distance = cumulative[segment]
        weight = 0.0 if end_distance == start_distance else (
            (target - start_distance) / (end_distance - start_distance)
        )
        start, end = path[segment - 1], path[segment]
        sampled.append([
            start[0] + (end[0] - start[0]) * weight,
            start[1] + (end[1] - start[1]) * weight,
        ])
    sampled[0] = list(path[0])
    sampled[-1] = list(path[-1])
    return sampled


def _render_filters(analysis: AnalysisData) -> dict[str, object]:
    selections: dict[str, object] = {}
    date_values = analysis.data["__data_inicio"].dropna()
    if not date_values.empty:
        minimum, maximum = date_values.min().date(), date_values.max().date()
        selected = st.date_input("Período", value=(minimum, maximum), min_value=minimum, max_value=maximum)
        if isinstance(selected, tuple) and len(selected) == 2:
            selections["periodo"] = selected

    labels = {
        "usuario": "Usuário", "unidade": "Unidade", "regiao": "Região",
        "motivo": "Motivo", "status": "Status",
    }
    available = [(key, analysis.columns[key]) for key in DASHBOARD_FILTER_FIELDS if key in analysis.columns]
    columns = st.columns(3)
    for position, (key, column) in enumerate(available):
        values = sorted(analysis.data[column].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist())
        selections[key] = columns[position % 3].multiselect(
            labels[key], values, key=f"filter_{key}", placeholder="Selecione uma ou mais opções"
        )
    if st.button("Limpar filtros"):
        for key in DASHBOARD_FILTER_FIELDS:
            st.session_state.pop(f"filter_{key}", None)
        st.rerun()
    return selections


def _format_metric(label: str, value: object) -> str:
    if value is None or pd.isna(value):
        return "Não disponível"
    if "quilômetro" in label.lower() or "média" in label.lower():
        return f"{float(value):,.2f} km".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{int(value):,}".replace(",", ".")


def _render_metrics(metrics: dict[str, object]) -> None:
    st.subheader("Indicadores")
    items = list(metrics.items())
    for start in range(0, len(items), 4):
        columns = st.columns(4)
        for column, (label, value) in zip(columns, items[start:start + 4]):
            column.metric(label, _format_metric(label, value))


def _chart(frame: pd.DataFrame, title: str, kind: str = "bar", x: str | None = None, y: str | None = None, color: str | None = None, light_mode: bool = False) -> None:
    if frame.empty or (y and frame[y].dropna().empty):
        st.info(f"Dados insuficientes para: {title}.")
        return
    if kind == "line":
        figure = px.line(frame, x=x, y=y, markers=True, title=title, color_discrete_sequence=[COLOR])
    elif kind == "pie":
        figure = px.pie(frame, names=x, values=y, hole=.45, title=title, color_discrete_sequence=px.colors.sequential.Blues_r)
    elif kind == "scatter":
        figure = px.scatter(frame, x=x, y=y, title=title, color=color, color_discrete_sequence=[COLOR], opacity=.72)
    else:
        figure = px.bar(frame, x=x, y=y, title=title, color_discrete_sequence=[COLOR])
    figure.update_layout(template="plotly_white" if light_mode else "plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", margin=dict(l=10, r=10, t=55, b=10))
    st.plotly_chart(figure, width="stretch")


def _ranking(routes: pd.DataFrame, category: str, metric: str, label: str, top: object) -> pd.DataFrame:
    frame = routes.dropna(subset=[category]).groupby(category, as_index=False)[metric].agg("count" if metric == category else "sum")
    frame = frame.rename(columns={category: label, metric: "Valor"}).sort_values("Valor", ascending=False)
    return frame if top == "Todos" else frame.head(int(top))


def _render_charts(analysis: AnalysisData, routes: pd.DataFrame, top: object, light_mode: bool = False) -> None:
    st.subheader("Visão geral")
    user, unit = analysis.columns.get("usuario"), analysis.columns.get("unidade")
    left, right = st.columns(2)
    with left:
        if user:
            counts = routes.groupby(user, as_index=False).size().rename(columns={user: "Usuário", "size": "Rotas"}).sort_values("Rotas", ascending=False)
            _chart(counts if top == "Todos" else counts.head(int(top)), "Pessoas com mais rotas", x="Usuário", y="Rotas", light_mode=light_mode)
        else:
            st.info("Coluna de usuário não identificada para o ranking de rotas.")
    with right:
        if user:
            km = routes.groupby(user, as_index=False)["__distancia_gps_valida"].sum().rename(columns={user: "Usuário", "__distancia_gps_valida": "Quilômetros"}).sort_values("Quilômetros", ascending=False)
            _chart(km if top == "Todos" else km.head(int(top)), "Pessoas com maior quilometragem", x="Usuário", y="Quilômetros", light_mode=light_mode)
        else:
            st.info("Coluna de usuário não identificada para o ranking de quilometragem.")

    dated = routes.dropna(subset=["__data_inicio"]).copy()
    if not dated.empty:
        dated["Data"] = dated["__data_inicio"].dt.date
        dated["Dia da semana"] = dated["__data_inicio"].dt.dayofweek.map(WEEKDAYS)
        dated["Hora"] = dated["__data_inicio"].dt.hour
    chart_pairs = st.columns(2)
    with chart_pairs[0]:
        daily = dated.groupby("Data", as_index=False).size().rename(columns={"size": "Rotas"}) if not dated.empty else pd.DataFrame()
        _chart(daily, "Quantidade de rotas por dia", "line", "Data", "Rotas", light_mode=light_mode)
    with chart_pairs[1]:
        if unit:
            by_unit = routes.groupby(unit, as_index=False)["__distancia_gps_valida"].sum().rename(columns={unit: "Unidade", "__distancia_gps_valida": "Quilômetros"})
            _chart(by_unit.sort_values("Quilômetros", ascending=False), "Quilometragem por unidade", x="Unidade", y="Quilômetros", light_mode=light_mode)
        else:
            st.info("Coluna de unidade não identificada.")

    distributions = st.columns(2)
    for container, (key, title) in zip(distributions, (("motivo", "Distribuição das rotas por motivo"), ("regiao", "Distribuição das rotas por região"))):
        with container:
            column = analysis.columns.get(key)
            if column:
                distribution = routes.groupby(column, as_index=False).size().rename(columns={column: "Categoria", "size": "Rotas"})
                _chart(distribution, title, "pie", "Categoria", "Rotas", light_mode=light_mode)
            else:
                st.info(f"Coluna de {key.replace('_', ' ')} não identificada.")

    temporal = st.columns(2)
    with temporal[0]:
        weekdays = dated.groupby("Dia da semana", as_index=False).size().rename(columns={"size": "Rotas"}) if not dated.empty else pd.DataFrame()
        _chart(weekdays, "Rotas por dia da semana", x="Dia da semana", y="Rotas", light_mode=light_mode)
    with temporal[1]:
        hourly = dated.groupby("Hora", as_index=False).size().rename(columns={"size": "Rotas"}) if not dated.empty else pd.DataFrame()
        _chart(hourly, "Rotas por horário de início", x="Hora", y="Rotas", light_mode=light_mode)

    comparison = routes.dropna(subset=["__distancia_gps_valida", "__distancia_hodometro"]).copy()
    _chart(comparison, "Distância GPS x hodômetro", "scatter", "__distancia_gps_valida", "__distancia_hodometro", light_mode=light_mode)


def _render_data_table(analysis: AnalysisData, filtered: pd.DataFrame) -> None:
    st.subheader("Dados das rotas")
    st.caption("Consulte os registros que correspondem aos filtros aplicados acima.")
    preferred = [
        "usuario", "unidade", "regiao", "motivo", "status", "data_inicio", "data_fim",
        "distancia_hodometro_existente", "distancia_gps",
    ]
    columns = [analysis.columns[key] for key in preferred if key in analysis.columns]
    by_normalized = {normalize_header(column): column for column in filtered.columns}
    percent_column = by_normalized.get("PERCENTUAL DIFERENCA %")
    link_column = by_normalized.get("LINK DA ROTA")
    columns.extend(column for column in (percent_column, link_column) if column and column not in columns)
    if not columns:
        columns = [column for column in filtered.columns if not str(column).startswith("__")]
    visible = filtered[columns].copy()
    labels = {
        analysis.columns.get("usuario"): "Nome",
        analysis.columns.get("unidade"): "Unidade",
        analysis.columns.get("regiao"): "Região",
        analysis.columns.get("motivo"): "Motivo",
        analysis.columns.get("status"): "Status",
        analysis.columns.get("data_inicio"): "Início",
        analysis.columns.get("data_fim"): "Fim",
        analysis.columns.get("distancia_hodometro_existente"): "Distância hodômetro (km)",
        analysis.columns.get("distancia_gps"): "Distância GPS (km)",
        percent_column: "Diferença (%)",
        link_column: "Rota",
    }
    visible = visible.rename(columns={key: value for key, value in labels.items() if key})
    if "Diferença (%)" in visible:
        visible["Diferença (%)"] = pd.to_numeric(visible["Diferença (%)"], errors="coerce") * 100
    column_config = {}
    if "Distância hodômetro (km)" in visible:
        column_config["Distância hodômetro (km)"] = st.column_config.NumberColumn(format="%.2f km")
    if "Distância GPS (km)" in visible:
        column_config["Distância GPS (km)"] = st.column_config.NumberColumn(format="%.2f km")
    if "Diferença (%)" in visible:
        column_config["Diferença (%)"] = st.column_config.NumberColumn(format="%.2f%%")
    if "Rota" in visible:
        column_config["Rota"] = st.column_config.LinkColumn(display_text="Abrir rota ↗")
    st.dataframe(visible, width="stretch", hide_index=True, height=460, column_config=column_config)


def _render_audit(audit: pd.DataFrame) -> None:
    st.subheader("Auditoria dos dados")
    if audit.empty:
        st.success("Nenhuma inconsistência foi identificada com os limites atuais.")
        return
    visible = audit.drop(columns=["Índice da linha"], errors="ignore")
    st.dataframe(visible, width="stretch", hide_index=True)
    st.download_button("Baixar inconsistências em CSV", visible.to_csv(index=False).encode("utf-8-sig"), "inconsistencias.csv", "text/csv")


def _render_downloads(analysis: AnalysisData, filtered: pd.DataFrame, audit: pd.DataFrame, metrics: dict[str, object], original_name: str) -> None:
    st.subheader("Baixar análise")
    st.caption("Os arquivos abaixo respeitam todos os filtros selecionados.")
    public_columns = [column for column in filtered.columns if not str(column).startswith("__")]
    stem = original_name.rsplit(".", 1)[0]
    left, middle, right = st.columns(3)
    left.download_button(
        "Baixar dados filtrados em CSV",
        filtered[public_columns].to_csv(index=False).encode("utf-8-sig"),
        f"{stem}_dados_filtrados.csv",
        "text/csv",
        width="stretch",
    )
    try:
        pdf_data = analysis_pdf(metrics, filtered, audit, analysis, original_name)
        middle.download_button(
            "Baixar relatório em PDF",
            pdf_data,
            f"{stem}_relatorio_analise.pdf",
            "application/pdf",
            width="stretch",
        )
    except Exception:
        logging.exception("Falha ao gerar o relatório PDF")
        middle.error("Não foi possível preparar o PDF.")
    try:
        excel_data = analysis_workbook(metrics, filtered, audit, analysis)
        right.download_button(
            "Baixar relatório em Excel",
            excel_data,
            f"{stem}_resumo_analise.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
        )
    except Exception:
        logging.exception("Falha ao gerar o relatório Excel")
        right.error("Não foi possível preparar o Excel.")
