from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from checklist_api import ChecklistSnapshot
from checklist_transform import ERP_COLUMNS, build_standardized_workbook, transformar_checklists_em_planilha
from route_processor import RouteResult, process_workbook


ITEMS = [
    {"itemId": 101, "categoryId": 2879584, "name": "Valor hodômetro inicial"},
    {"itemId": 102, "categoryId": 2879584, "name": "Valor hodômetro final"},
    {"itemId": 103, "categoryId": 2879584, "name": "Motivo"},
    {"itemId": 104, "categoryId": 2879584, "name": "Centro de custo"},
]


def snapshot(*, answer_category=2879584, duplicate_count=0):
    evaluation = {
        "evaluationId": 55,
        "checklistId": 9,
        "status": 6,
        "userId": 3,
        "unitId": 4,
        "startedAt": "2026-10-05T11:00:00+00:00",
        "concludedAt": "2026-10-05T12:00:00+00:00",
        "initialLatitude": -23.5,
        "initialLongitude": -46.6,
        "finalLatitude": -22.9,
        "finalLongitude": -43.2,
    }
    answers = [
        {"resultId": 1, "itemId": 101, "categoryId": answer_category, "number": 1000, "answeredAt": "2026-10-05T11:01:00+00:00"},
        {"resultId": 2, "itemId": 102, "categoryId": answer_category, "number": 1125.5, "answeredAt": "2026-10-05T11:59:00+00:00"},
        {"resultId": 3, "itemId": 103, "categoryId": answer_category, "selectedOptions": [{"optionText": "Visita"}], "answeredAt": "2026-10-05T11:30:00+00:00"},
        {"resultId": 4, "itemId": 104, "categoryId": answer_category, "text": "CC-100", "answeredAt": "2026-10-05T11:30:00+00:00"},
    ]
    return ChecklistSnapshot(
        category_id=2879584,
        checklist_id=9,
        year=2026,
        month=10,
        evaluations=[evaluation],
        results_by_evaluation={55: answers},
        users=[{"userId": 3, "name": "Ana", "userTypeId": 7}],
        user_types=[{"userTypeId": 7, "name": "Motorista"}],
        units=[{"unitId": 4, "name": "Unidade SP"}],
        unit_regions=[{"unitId": 4, "regionId": 8}],
        regions=[{"regionId": 8, "name": "Sudeste", "hierarchyName": "Brasil > Sudeste"}],
        items=ITEMS,
        duplicate_count=duplicate_count,
    )


def fake_route(origin, destination):
    return RouteResult(321.45, "Calculada", "https://www.google.com/maps/dir/?api=1")


def manual_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COORDENADA GPS INICIAL", "COORDENADA GPS FINAL"])
    sheet.append(["-23.5, -46.6", "-22.9, -43.2"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_transforms_api_data_to_erp_columns_and_local_dates():
    result = transformar_checklists_em_planilha(snapshot(duplicate_count=2))
    assert result.found == result.valid == 1
    assert result.inconsistent == 0
    assert result.duplicate_count == 2
    assert list(result.valid_rows.columns) == ERP_COLUMNS
    row = result.valid_rows.iloc[0]
    assert row["ID RESPOSTA"] == 55
    assert row["DATA INÍCIO"].strftime("%d/%m/%Y %H:%M") == "05/10/2026 08:00"
    assert row["DISTÂNCIA HODÔMETRO"] == 125.5
    assert row["COORDENADA GPS INICIAL"] == "-23.5, -46.6"
    assert row["MOTIVO"] == "Visita"


def test_odometer_mapping_does_not_confuse_gps_items_with_vehicle_odometer():
    data = snapshot()
    data.items[0]["name"] = "Insira a Quilometragem inicial do seu veículo (hodômetro)"
    data.items[1]["name"] = "Insira a Quilometragem Final do seu veículo (hodômetro)"
    data.items.extend(
        [
            {"itemId": 201, "categoryId": 2879584, "name": "Quilometragem inicial (GPS)"},
            {"itemId": 202, "categoryId": 2879584, "name": "Quilometragem final (GPS)"},
        ]
    )
    data.results_by_evaluation[55].extend(
        [
            {
                "resultId": 5,
                "itemId": 201,
                "categoryId": 2879584,
                "text": "-23.5,-46.6",
                "answeredAt": "2026-10-05T11:01:00+00:00",
            },
            {
                "resultId": 6,
                "itemId": 202,
                "categoryId": 2879584,
                "text": "-22.9,-43.2",
                "answeredAt": "2026-10-05T11:59:00+00:00",
            },
        ]
    )

    result = transformar_checklists_em_planilha(data)
    row = result.valid_rows.iloc[0]

    assert result.resolved_mapping["odometer_start_item_id"] == 101
    assert result.resolved_mapping["odometer_end_item_id"] == 102
    assert row["VALOR ODOMETRO INICIAL"] == 1000
    assert row["VALOR ODOMETRO FINAL"] == 1125.5
    assert row["DISTÂNCIA HODÔMETRO"] == 125.5


def test_optional_answers_from_another_category_do_not_block_route_calculation():
    result = transformar_checklists_em_planilha(snapshot(answer_category=999))
    assert result.valid == 1
    assert result.inconsistent == 0
    assert result.valid_rows.iloc[0]["VALIDAÇÃO"] == "Válida"
    assert result.valid_rows.iloc[0]["VALOR ODOMETRO INICIAL"] is None


def test_missing_route_coordinate_is_the_only_reason_to_block_calculation():
    data = snapshot()
    data.evaluations[0]["initialLatitude"] = None
    data.evaluations[0]["initialLongitude"] = None

    result = transformar_checklists_em_planilha(data)

    assert result.valid == 0
    assert result.inconsistent == 1
    assert "coordenada inicial" in result.inconsistent_rows.iloc[0]["VALIDAÇÃO"]


def test_generated_workbook_matches_erp_layout_and_separates_inconsistencies():
    contract_path = Path(__file__).parent / "fixtures" / "erp_contract.json"
    contract = json.loads(contract_path.read_text(encoding="utf-8"))
    valid_result = transformar_checklists_em_planilha(snapshot())
    workbook = load_workbook(BytesIO(build_standardized_workbook(valid_result)))
    assert workbook.sheetnames == [contract["sheet_name"]]
    assert [cell.value for cell in workbook.active[1]] == contract["columns"] == ERP_COLUMNS
    assert workbook.active["G2"].number_format == "dd/mm/yyyy hh:mm:ss"

    invalid_snapshot = snapshot()
    invalid_snapshot.evaluations[0]["finalLatitude"] = None
    invalid_snapshot.evaluations[0]["finalLongitude"] = None
    invalid_result = transformar_checklists_em_planilha(invalid_snapshot)
    invalid_workbook = load_workbook(BytesIO(build_standardized_workbook(invalid_result)))
    assert invalid_workbook.sheetnames == ["CONTROLE QUILOMETRAGEM", "INCONSISTENCIAS"]
    assert invalid_workbook["CONTROLE QUILOMETRAGEM"].max_row == 1
    assert invalid_workbook["INCONSISTENCIAS"].max_row == 2


def test_existing_route_calculation_produces_same_result_for_manual_and_api_sources():
    api_workbook = build_standardized_workbook(transformar_checklists_em_planilha(snapshot()))
    manual = process_workbook(manual_workbook(), route_function=fake_route)
    imported = process_workbook(api_workbook, route_function=fake_route)
    assert (manual.total, manual.calculated, manual.failed) == (1, 1, 0)
    assert (imported.total, imported.calculated, imported.failed) == (1, 1, 0)

    manual_sheet = load_workbook(BytesIO(manual.workbook)).active
    imported_sheet = load_workbook(BytesIO(imported.workbook)).active
    manual_headers = {cell.value: cell.column for cell in manual_sheet[1]}
    imported_headers = {cell.value: cell.column for cell in imported_sheet[1]}
    assert manual_sheet.cell(2, manual_headers["DISTÂNCIA GPS"]).value == imported_sheet.cell(
        2, imported_headers["DISTÂNCIA GPS"]
    ).value
    assert manual_sheet.cell(2, manual_headers["STATUS DA ROTA"]).value == "Calculada"
    assert imported_sheet.cell(2, imported_headers["STATUS DA ROTA"]).value == "Calculada"
    assert imported_sheet.cell(2, imported_headers["DISTÂNCIA HODÔMETRO"]).value == 125.5
    expected_percentage = abs(125.5 - 321.45) / 321.45
    percentage_cell = imported_sheet.cell(2, imported_headers["PERCENTUAL DIFERENÇA %"])
    assert percentage_cell.value == expected_percentage
    assert percentage_cell.number_format == "0.00%"
