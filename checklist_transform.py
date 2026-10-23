"""Transforma dados da API Analytics no padrão de quilometragem do ERP."""

from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any, Iterable, Mapping
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from checklist_api import ChecklistSnapshot, DEFAULT_TIMEZONE, index_by_id, parse_api_datetime

SHEET_NAME = "CONTROLE QUILOMETRAGEM"
INCONSISTENCY_SHEET = "INCONSISTENCIAS"

ERP_COLUMNS = [
    "DESCRIÇÃO DO STATUS",
    "ID RESPOSTA",
    "NOME",
    "TIPO DE USUÁRIO",
    "UNIDADE",
    "REGIÃO",
    "DATA INÍCIO",
    "DATA CONCLUSÃO",
    "COORDENADA GPS INICIAL",
    "VALOR ODOMETRO INICIAL",
    "COORDENADA GPS FINAL",
    "VALOR ODOMETRO FINAL",
    "MOTIVO",
    "CENTRO DE CUSTO",
    "DISTÂNCIA HODÔMETRO",
    "VALIDAÇÃO",
    "PERCENTUAL DIFERENÇA %",
    "DISTÂNCIA GPS",
    "STATUS DA ROTA",
    "LINK DA ROTA",
]

MAPPING_ALIASES = {
    "odometer_start_item_id": (
        "VALOR ODOMETRO INICIAL", "ODOMETRO INICIAL", "HODOMETRO INICIAL",
        "QUILOMETRAGEM INICIAL", "KM INICIAL",
    ),
    "odometer_end_item_id": (
        "VALOR ODOMETRO FINAL", "ODOMETRO FINAL", "HODOMETRO FINAL",
        "QUILOMETRAGEM FINAL", "KM FINAL",
    ),
    "reason_item_id": ("MOTIVO", "MOTIVO DA ROTA", "FINALIDADE", "TIPO DE DESLOCAMENTO"),
    "cost_center_item_id": ("CENTRO DE CUSTO", "CENTRO CUSTO"),
    "gps_start_item_id": ("COORDENADA GPS INICIAL", "GPS INICIAL", "LOCALIZACAO INICIAL"),
    "gps_end_item_id": ("COORDENADA GPS FINAL", "GPS FINAL", "LOCALIZACAO FINAL"),
}

# Nenhuma pergunta complementar é obrigatória para calcular uma rota. O
# processamento depende somente das coordenadas inicial e final, que também
# podem vir diretamente dos campos de geolocalização da avaliação.
REQUIRED_MAPPINGS: tuple[str, ...] = ()


class ChecklistTransformationError(ValueError):
    """Configuração ou dados incompatíveis com a transformação."""


@dataclass(frozen=True)
class TransformationResult:
    all_rows: pd.DataFrame
    valid_rows: pd.DataFrame
    inconsistent_rows: pd.DataFrame
    resolved_mapping: dict[str, int | None]
    mapping_warnings: tuple[str, ...]
    duplicate_count: int

    @property
    def found(self) -> int:
        return len(self.all_rows)

    @property
    def valid(self) -> int:
        return len(self.valid_rows)

    @property
    def inconsistent(self) -> int:
        return len(self.inconsistent_rows)


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().upper()


def _optional_int(value: object) -> int | None:
    if value in (None, "", 0, "0"):
        return None
    try:
        return int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError) as exc:
        raise ChecklistTransformationError(f"O ID de item {value!r} não é um número válido.") from exc


def resolve_item_mapping(
    items: Iterable[dict[str, Any]],
    configured: Mapping[str, object] | None,
    category_id: int,
) -> tuple[dict[str, int | None], tuple[str, ...]]:
    """Valida IDs configurados e descobre IDs por nomes inequívocos quando possível."""

    category_items = [row for row in items if _optional_int(row.get("categoryId")) == category_id]
    by_id = {
        int(row["itemId"]): row
        for row in category_items
        if _optional_int(row.get("itemId")) is not None
    }
    configured = configured or {}
    resolved: dict[str, int | None] = {}
    warnings: list[str] = []
    for field, aliases in MAPPING_ALIASES.items():
        explicit = _optional_int(configured.get(field))
        if explicit is not None:
            if explicit not in by_id:
                raise ChecklistTransformationError(
                    f"O item {explicit}, configurado para {field}, não pertence à categoria {category_id}."
                )
            resolved[field] = explicit
            continue

        normalized_aliases = tuple(normalize_text(alias) for alias in aliases)
        exact = [
            item_id for item_id, row in by_id.items()
            if normalize_text(row.get("name")) in normalized_aliases
        ]
        candidates = exact
        if not candidates:
            candidates = [
                item_id for item_id, row in by_id.items()
                if any(alias in normalize_text(row.get("name")) for alias in normalized_aliases)
            ]

        # "Quilometragem inicial/final (GPS)" contém a palavra
        # quilometragem, mas seu valor é uma coordenada. Para os campos de
        # hodômetro, prioriza perguntas que declaram explicitamente
        # HODÔMETRO/ODÔMETRO e exclui itens GPS do fallback.
        if field.startswith("odometer_") and candidates:
            odometer_candidates = [
                item_id
                for item_id in candidates
                if any(
                    term in normalize_text(by_id[item_id].get("name"))
                    for term in ("HODOMETRO", "ODOMETRO")
                )
            ]
            if odometer_candidates:
                candidates = odometer_candidates
            else:
                candidates = [
                    item_id
                    for item_id in candidates
                    if "GPS" not in normalize_text(by_id[item_id].get("name"))
                ]
        if len(candidates) == 1:
            resolved[field] = candidates[0]
        else:
            resolved[field] = None
            label = field.replace("_item_id", "").replace("_", " ")
            if len(candidates) > 1 and field in REQUIRED_MAPPINGS:
                warnings.append(f"Mais de um item pode representar {label}; configure o ID explicitamente.")
            elif field in REQUIRED_MAPPINGS:
                warnings.append(f"Não foi possível identificar automaticamente o item de {label}.")
    return resolved, tuple(warnings)


def _answer_value(answer: dict[str, Any]) -> object:
    if answer.get("number") is not None:
        return answer["number"]
    text = answer.get("text")
    if text is not None and str(text).strip() != "":
        return str(text).replace("<br />", "\n").strip()
    options = answer.get("selectedOptions")
    if isinstance(options, list):
        values = []
        for option in options:
            if isinstance(option, dict):
                value = option.get("optionText")
            else:
                value = option
            if value is not None and str(value).strip():
                values.append(str(value).strip())
        return "; ".join(dict.fromkeys(values)) if values else None
    records = answer.get("records")
    if isinstance(records, dict):
        return records.get("name") or records.get("reference")
    return None


def _answers_for_item(answers: list[dict[str, Any]], item_id: int | None) -> list[dict[str, Any]]:
    if item_id is None:
        return []
    return [answer for answer in answers if _optional_int(answer.get("itemId")) == item_id]


def _answer_for_field(
    answers: list[dict[str, Any]], item_id: int | None, *, prefer_earliest: bool = False
) -> object:
    matches = _answers_for_item(answers, item_id)
    if not matches:
        return None
    minimum = datetime.min.replace(tzinfo=ZoneInfo("UTC"))
    matches.sort(
        key=lambda row: (
            parse_api_datetime(row.get("answeredAt")) or minimum,
            _optional_int(row.get("resultId")) or 0,
        ),
        reverse=not prefer_earliest,
    )
    return _answer_value(matches[0])


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = re.sub(r"[^\d,.-]", "", str(value).strip())
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _coordinate(value: object) -> str | None:
    if value is None:
        return None
    parts = [part.strip() for part in str(value).split(",")]
    if len(parts) != 2:
        return None
    latitude, longitude = _number(parts[0]), _number(parts[1])
    if latitude is None or longitude is None or not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
        return None
    latitude_text = f"{latitude:.7f}".rstrip("0").rstrip(".")
    longitude_text = f"{longitude:.7f}".rstrip("0").rstrip(".")
    return f"{latitude_text}, {longitude_text}"


def _coordinate_from_evaluation(evaluation: dict[str, Any], prefix: str) -> str | None:
    latitude = _number(evaluation.get(f"{prefix}Latitude"))
    longitude = _number(evaluation.get(f"{prefix}Longitude"))
    if latitude is None or longitude is None:
        return None
    return _coordinate(f"{latitude},{longitude}")


def _local_datetime(value: object, timezone_name: str) -> datetime | None:
    parsed = parse_api_datetime(value)
    if parsed is None:
        return None
    return parsed.astimezone(ZoneInfo(timezone_name)).replace(tzinfo=None)


def _nonempty(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def transformar_checklists_em_planilha(
    snapshot: ChecklistSnapshot,
    item_mapping: Mapping[str, object] | None = None,
    *,
    timezone_name: str = DEFAULT_TIMEZONE,
) -> TransformationResult:
    """Converte o snapshot em linhas válidas e inconsistentes no padrão ERP."""

    mapping, warnings = resolve_item_mapping(snapshot.items, item_mapping, snapshot.category_id)
    users = index_by_id(snapshot.users, "userId")
    user_types = index_by_id(snapshot.user_types, "userTypeId")
    units = index_by_id(snapshot.units, "unitId")
    regions = index_by_id(snapshot.regions, "regionId")
    region_ids_by_unit: dict[int, list[int]] = {}
    for relation in snapshot.unit_regions:
        unit_id = _optional_int(relation.get("unitId"))
        region_id = _optional_int(relation.get("regionId"))
        if unit_id is not None and region_id is not None:
            region_ids_by_unit.setdefault(unit_id, []).append(region_id)

    records: list[dict[str, object]] = []
    for evaluation in snapshot.evaluations:
        evaluation_id = _optional_int(evaluation.get("evaluationId"))
        if evaluation_id is None:
            continue
        answers = [
            answer for answer in snapshot.results_by_evaluation.get(evaluation_id, [])
            if _optional_int(answer.get("categoryId")) == snapshot.category_id
        ]
        user_id = _optional_int(evaluation.get("userId"))
        unit_id = _optional_int(evaluation.get("unitId"))
        user = users.get(user_id or -1, {})
        unit = units.get(unit_id or -1, {})
        user_type = user_types.get(_optional_int(user.get("userTypeId")) or -1, {})
        region_names = []
        for region_id in region_ids_by_unit.get(unit_id or -1, []):
            region = regions.get(region_id, {})
            name = region.get("hierarchyName") or region.get("name")
            if _nonempty(name):
                region_names.append(str(name).strip())

        gps_start = _coordinate(_answer_for_field(answers, mapping.get("gps_start_item_id"), prefer_earliest=True))
        gps_end = _coordinate(_answer_for_field(answers, mapping.get("gps_end_item_id")))
        gps_start = gps_start or _coordinate_from_evaluation(evaluation, "initial")
        gps_end = gps_end or _coordinate_from_evaluation(evaluation, "final")
        odometer_start = _number(
            _answer_for_field(answers, mapping.get("odometer_start_item_id"), prefer_earliest=True)
        )
        odometer_end = _number(_answer_for_field(answers, mapping.get("odometer_end_item_id")))
        reason = _answer_for_field(answers, mapping.get("reason_item_id"))
        cost_center = _answer_for_field(answers, mapping.get("cost_center_item_id"))
        started_at = _local_datetime(evaluation.get("startedAt"), timezone_name)
        concluded_at = _local_datetime(evaluation.get("concludedAt"), timezone_name)

        # Somente as coordenadas determinam se uma linha pode seguir para o
        # roteador. Dados cadastrais, hodômetro, motivo e centro de custo são
        # mantidos quando existirem, mas nunca bloqueiam o cálculo GPS.
        route_problems: list[str] = []
        required_route_values = {
            "coordenada inicial": gps_start,
            "coordenada final": gps_end,
        }
        route_problems.extend(
            label for label, value in required_route_values.items() if not _nonempty(value)
        )
        distance = None
        if odometer_start is not None and odometer_end is not None:
            distance = odometer_end - odometer_start

        validation = (
            "Válida"
            if not route_problems
            else "Sem rota: " + "; ".join(dict.fromkeys(route_problems))
        )
        records.append(
            {
                "DESCRIÇÃO DO STATUS": "Concluído",
                "ID RESPOSTA": evaluation_id,
                "NOME": user.get("name"),
                "TIPO DE USUÁRIO": user_type.get("name"),
                "UNIDADE": unit.get("name"),
                "REGIÃO": "; ".join(dict.fromkeys(region_names)) or None,
                "DATA INÍCIO": started_at,
                "DATA CONCLUSÃO": concluded_at,
                "COORDENADA GPS INICIAL": gps_start,
                "VALOR ODOMETRO INICIAL": odometer_start,
                "COORDENADA GPS FINAL": gps_end,
                "VALOR ODOMETRO FINAL": odometer_end,
                "MOTIVO": reason,
                "CENTRO DE CUSTO": cost_center,
                "DISTÂNCIA HODÔMETRO": distance,
                "VALIDAÇÃO": validation,
                "PERCENTUAL DIFERENÇA %": None,
                "DISTÂNCIA GPS": None,
                "STATUS DA ROTA": None,
                "LINK DA ROTA": None,
            }
        )

    all_rows = pd.DataFrame(records, columns=ERP_COLUMNS)
    if all_rows.empty:
        valid_rows = all_rows.copy()
        inconsistent_rows = all_rows.copy()
    else:
        valid_mask = all_rows["VALIDAÇÃO"].eq("Válida")
        valid_rows = all_rows.loc[valid_mask].reset_index(drop=True)
        inconsistent_rows = all_rows.loc[~valid_mask].reset_index(drop=True)
    return TransformationResult(
        all_rows=all_rows,
        valid_rows=valid_rows,
        inconsistent_rows=inconsistent_rows,
        resolved_mapping=mapping,
        mapping_warnings=warnings,
        duplicate_count=snapshot.duplicate_count,
    )


def _safe_excel_value(value: object) -> object:
    if pd.isna(value):
        return None
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def _write_dataframe(worksheet, dataframe: pd.DataFrame) -> None:
    worksheet.append(list(dataframe.columns))
    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append([_safe_excel_value(value) for value in row])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 30
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 24, "B": 16, "C": 28, "D": 22, "E": 24, "F": 28,
        "G": 21, "H": 21, "I": 28, "J": 24, "K": 28, "L": 24,
        "M": 28, "N": 22, "O": 24, "P": 52, "Q": 25, "R": 20,
        "S": 28, "T": 18,
    }
    for column, width in widths.items():
        if worksheet.max_column >= ord(column) - 64:
            worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (7, 8) and isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            elif cell.column in (10, 12, 15, 18) and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
            elif cell.column == 17 and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def build_standardized_workbook(result: TransformationResult) -> bytes:
    """Cria o Excel padronizado; inconsistências ficam separadas da base calculável."""

    workbook = Workbook()
    valid_sheet = workbook.active
    valid_sheet.title = SHEET_NAME
    _write_dataframe(valid_sheet, result.valid_rows)
    if not result.inconsistent_rows.empty:
        invalid_sheet = workbook.create_sheet(INCONSISTENCY_SHEET)
        _write_dataframe(invalid_sheet, result.inconsistent_rows)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def item_catalog_dataframe(snapshot: ChecklistSnapshot) -> pd.DataFrame:
    """Tabela pequena para ajudar a configurar os IDs sem expor respostas."""

    rows = [
        {
            "ID do item": _optional_int(item.get("itemId")),
            "Nome do item": item.get("name"),
            "ID da categoria": _optional_int(item.get("categoryId")),
        }
        for item in snapshot.items
    ]
    return pd.DataFrame(rows, columns=["ID do item", "Nome do item", "ID da categoria"])
